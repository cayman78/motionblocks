#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
find_site_plans.py

Ищет в большой папке файлы/страницы, похожие на:
- план участка
- генплан
- кадастровый / межевой план
- ситуационный план
- топографический план
- схему расположения дома на участке

Что делает:
1. Рекурсивно обходит папку.
2. Обрабатывает PDF, JPG/PNG/TIFF/BMP/WEBP, DOCX, XLSX, ZIP.
3. По PDF рендерит страницы в картинки и делает OCR, если установлен Tesseract.
4. По изображениям делает OCR.
5. По DOCX/XLSX извлекает текст.
6. Считает score по ключевым словам.
7. Складывает кандидаты в output/candidates.
8. Пишет report.csv.

Установка зависимостей:

    pip install pymupdf pillow pytesseract python-docx openpyxl tqdm

Для OCR нужно отдельно установить Tesseract OCR:
Windows:
    winget install UB-Mannheim.TesseractOCR

Потом, если нужно, указать путь:
    set TESSERACT_CMD=C:\Program Files\Tesseract-OCR\tesseract.exe

Пример запуска:

    python find_site_plans.py "D:\big_unsorted_folder" --out "D:\site_plan_search_result"

Более мягкий режим, чтобы поймать больше кандидатов:

    python find_site_plans.py "D:\big_unsorted_folder" --out "D:\site_plan_search_result" --min-score 2

Ограничить количество страниц PDF для теста:

    python find_site_plans.py "D:\big_unsorted_folder" --out "D:\site_plan_search_result" --pdf-max-pages 10
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import sys
import tempfile
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable, Optional

try:
    from PIL import Image, ImageOps
except ImportError:
    Image = None
    ImageOps = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from tqdm import tqdm
except ImportError:
    tqdm = None


IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
ZIP_EXTS = {".zip"}

# Слова и фразы. Вес выше — признак сильнее.
KEYWORDS = {
    # Очень сильные признаки
    "кадастровый план": 8,
    "межевой план": 8,
    "генеральный план": 8,
    "генплан": 8,
    "ситуационный план": 8,
    "топографический план": 8,
    "топосъемка": 8,
    "топосъёмка": 8,
    "план участка": 8,
    "схема расположения земельного участка": 9,
    "границы земельного участка": 9,

    # Сильные признаки
    "земельный участок": 5,
    "кадастровый номер": 5,
    "кадастровый квартал": 5,
    "схема расположения": 5,
    "пятно застройки": 5,
    "отступ от границы": 5,
    "граница участка": 5,
    "координаты характерных точек": 6,
    "характерные точки границ": 6,

    # Средние признаки
    "участок": 2,
    "план": 2,
    "граница": 2,
    "границы": 2,
    "застройка": 2,
    "координаты": 2,
    "экспликация": 2,
    "масштаб": 2,
    "м 1:": 3,
    "площадь участка": 4,

    # Английские варианты, если часть файлов от проектировщиков
    "site plan": 8,
    "plot plan": 8,
    "land plot": 6,
    "cadastral plan": 8,
    "topographic survey": 8,
    "boundary plan": 7,
    "property boundary": 6,
    "master plan": 5,
}

# Слова, которые часто встречаются на планах/чертежах, но не являются гарантией.
DRAWING_HINTS = {
    "оси": 1,
    "экспликация зданий": 4,
    "условные обозначения": 3,
    "существующее здание": 3,
    "проектируемое здание": 3,
    "красные линии": 4,
    "отмостка": 2,
    "забор": 2,
    "ворота": 2,
    "септик": 2,
    "скважина": 2,
    "пятно": 2,
    "инженерные сети": 3,
}


@dataclass
class Finding:
    source_file: str
    file_type: str
    page_or_item: str
    score: int
    matched_terms: str
    candidate_artifact: str
    notes: str


def normalize_text(text: str) -> str:
    text = text or ""
    text = text.replace("ё", "е").replace("Ё", "Е")
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def score_text(text: str, file_name: str = "") -> tuple[int, list[str]]:
    haystack = normalize_text(file_name + " " + text)

    score = 0
    matched: list[str] = []

    for term, weight in KEYWORDS.items():
        norm_term = normalize_text(term)
        if norm_term in haystack:
            score += weight
            matched.append(term)

    for term, weight in DRAWING_HINTS.items():
        norm_term = normalize_text(term)
        if norm_term in haystack:
            score += weight
            matched.append(term)

    # Регулярки для кадастрового номера: 50:20:0010203:123 и похожие.
    cadastral_patterns = [
        r"\b\d{2}:\d{2}:\d{6,7}:\d+\b",
        r"\b\d{2}:\d{2}:\d{6,7}\b",
    ]
    for pat in cadastral_patterns:
        if re.search(pat, haystack):
            score += 6
            matched.append("regex: cadastral_number")

    # Координатные таблицы / точки н1, н2...
    if re.search(r"\b[нn]\s*[\d]{1,3}\b", haystack) and "координат" in haystack:
        score += 4
        matched.append("regex: boundary_points")

    return score, sorted(set(matched))


def ensure_tesseract_configured() -> bool:
    if pytesseract is None:
        return False

    cmd = os.environ.get("TESSERACT_CMD")
    if cmd:
        pytesseract.pytesseract.tesseract_cmd = cmd

    try:
        _ = pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def ocr_image(image_path: Path, lang: str) -> str:
    if Image is None or pytesseract is None:
        return ""

    try:
        img = Image.open(image_path)

        # Немного нормализуем картинку для OCR.
        img = ImageOps.exif_transpose(img)
        if img.mode not in ("L", "RGB"):
            img = img.convert("RGB")

        # Для больших чертежей OCR иногда лучше на сером.
        gray = img.convert("L")

        return pytesseract.image_to_string(gray, lang=lang)
    except Exception as e:
        return f""


def safe_copy(src: Path, dst_dir: Path, prefix: str = "") -> Path:
    dst_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-zА-Яа-я0-9_.\-]+", "_", src.name)
    dst = dst_dir / f"{prefix}{safe_name}"
    base = dst.stem
    ext = dst.suffix
    i = 1
    while dst.exists():
        dst = dst_dir / f"{base}_{i}{ext}"
        i += 1
    shutil.copy2(src, dst)
    return dst


def save_image_candidate(src_image: Path, out_candidates: Path, label: str) -> Path:
    return safe_copy(src_image, out_candidates, prefix=f"{label}__")


def iter_files(root: Path, include_archives: bool) -> Iterable[Path]:
    allowed = IMAGE_EXTS | PDF_EXTS | DOCX_EXTS | XLSX_EXTS
    if include_archives:
        allowed |= ZIP_EXTS

    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in allowed:
            yield path


def process_image(
    path: Path,
    out_candidates: Path,
    min_score: int,
    ocr_lang: str,
    tesseract_ok: bool,
) -> list[Finding]:
    text = ocr_image(path, ocr_lang) if tesseract_ok else ""
    score, matched = score_text(text, path.name)

    if score >= min_score:
        candidate = save_image_candidate(path, out_candidates, "image")
        return [
            Finding(
                source_file=str(path),
                file_type="image",
                page_or_item="image",
                score=score,
                matched_terms=", ".join(matched),
                candidate_artifact=str(candidate),
                notes="OCR used" if tesseract_ok else "OCR unavailable; matched by filename only",
            )
        ]
    return []


def render_pdf_page_to_image(pdf_path: Path, page_index: int, tmp_dir: Path, zoom: float = 1.7) -> Path:
    if fitz is None:
        raise RuntimeError("PyMuPDF is not installed. Run: pip install pymupdf")

    doc = fitz.open(pdf_path)
    try:
        page = doc.load_page(page_index)
        matrix = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        out = tmp_dir / f"{pdf_path.stem}_page_{page_index + 1}.png"
        pix.save(out)
        return out
    finally:
        doc.close()


def extract_pdf_text(pdf_path: Path, page_index: int) -> str:
    if fitz is None:
        return ""
    try:
        doc = fitz.open(pdf_path)
        try:
            page = doc.load_page(page_index)
            return page.get_text("text") or ""
        finally:
            doc.close()
    except Exception:
        return ""


def process_pdf(
    path: Path,
    out_candidates: Path,
    min_score: int,
    ocr_lang: str,
    tesseract_ok: bool,
    pdf_max_pages: Optional[int],
) -> list[Finding]:
    findings: list[Finding] = []

    if fitz is None:
        # Без PyMuPDF можем только скорить имя файла.
        score, matched = score_text("", path.name)
        if score >= min_score:
            candidate = safe_copy(path, out_candidates, "pdf_name_only__")
            findings.append(
                Finding(
                    source_file=str(path),
                    file_type="pdf",
                    page_or_item="unknown",
                    score=score,
                    matched_terms=", ".join(matched),
                    candidate_artifact=str(candidate),
                    notes="PyMuPDF unavailable; matched by filename only",
                )
            )
        return findings

    try:
        doc = fitz.open(path)
        n_pages = doc.page_count
        doc.close()
    except Exception as e:
        return [
            Finding(
                source_file=str(path),
                file_type="pdf",
                page_or_item="error",
                score=0,
                matched_terms="",
                candidate_artifact="",
                notes=f"Cannot open PDF: {e}",
            )
        ]

    pages_to_process = n_pages
    if pdf_max_pages is not None:
        pages_to_process = min(pages_to_process, pdf_max_pages)

    with tempfile.TemporaryDirectory() as td:
        tmp_dir = Path(td)

        for page_index in range(pages_to_process):
            page_label = f"page_{page_index + 1}"

            # Сначала текстовый слой PDF — быстро и дешево.
            text_layer = extract_pdf_text(path, page_index)

            # Потом OCR рендера страницы, если нужен.
            image_path = None
            ocr_text = ""

            # Рендерим страницу всегда, если score по тексту уже проходной,
            # чтобы сохранить картинку-кандидат для быстрой ручной проверки.
            score_pre, matched_pre = score_text(text_layer, path.name)

            need_ocr = tesseract_ok and score_pre < min_score
            need_preview = score_pre >= min_score

            if need_ocr or need_preview:
                try:
                    image_path = render_pdf_page_to_image(path, page_index, tmp_dir)
                except Exception:
                    image_path = None

            if need_ocr and image_path:
                ocr_text = ocr_image(image_path, ocr_lang)

            combined_text = f"{text_layer}\n{ocr_text}"
            score, matched = score_text(combined_text, path.name)

            if score >= min_score:
                if image_path is None:
                    try:
                        image_path = render_pdf_page_to_image(path, page_index, tmp_dir)
                    except Exception:
                        image_path = None

                if image_path:
                    candidate = save_image_candidate(
                        image_path,
                        out_candidates,
                        f"pdf__{path.stem}__p{page_index + 1}",
                    )
                else:
                    candidate = safe_copy(path, out_candidates, f"pdf__p{page_index + 1}__")

                findings.append(
                    Finding(
                        source_file=str(path),
                        file_type="pdf",
                        page_or_item=page_label,
                        score=score,
                        matched_terms=", ".join(matched),
                        candidate_artifact=str(candidate),
                        notes=(
                            f"PDF pages: {n_pages}; "
                            f"text_layer={'yes' if text_layer.strip() else 'no'}; "
                            f"OCR={'yes' if ocr_text.strip() else 'no'}"
                        ),
                    )
                )

    return findings


def process_docx(path: Path, out_candidates: Path, min_score: int) -> list[Finding]:
    if docx is None:
        score, matched = score_text("", path.name)
        if score >= min_score:
            candidate = safe_copy(path, out_candidates, "docx_name_only__")
            return [
                Finding(
                    source_file=str(path),
                    file_type="docx",
                    page_or_item="document",
                    score=score,
                    matched_terms=", ".join(matched),
                    candidate_artifact=str(candidate),
                    notes="python-docx unavailable; matched by filename only",
                )
            ]
        return []

    try:
        d = docx.Document(str(path))
        text = "\n".join(p.text for p in d.paragraphs)
        score, matched = score_text(text, path.name)

        if score >= min_score:
            candidate = safe_copy(path, out_candidates, "docx__")
            return [
                Finding(
                    source_file=str(path),
                    file_type="docx",
                    page_or_item="document_text",
                    score=score,
                    matched_terms=", ".join(matched),
                    candidate_artifact=str(candidate),
                    notes="DOCX text matched",
                )
            ]
    except Exception as e:
        return [
            Finding(
                source_file=str(path),
                file_type="docx",
                page_or_item="error",
                score=0,
                matched_terms="",
                candidate_artifact="",
                notes=f"Cannot read DOCX: {e}",
            )
        ]

    return []


def process_xlsx(path: Path, out_candidates: Path, min_score: int) -> list[Finding]:
    if openpyxl is None:
        score, matched = score_text("", path.name)
        if score >= min_score:
            candidate = safe_copy(path, out_candidates, "xlsx_name_only__")
            return [
                Finding(
                    source_file=str(path),
                    file_type="xlsx",
                    page_or_item="workbook",
                    score=score,
                    matched_terms=", ".join(matched),
                    candidate_artifact=str(candidate),
                    notes="openpyxl unavailable; matched by filename only",
                )
            ]
        return []

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        chunks: list[str] = []
        for ws in wb.worksheets:
            chunks.append(ws.title)
            # Чтобы не читать гигантские книги бесконечно: первые 200 строк на лист.
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    break
                for val in row:
                    if val is not None:
                        chunks.append(str(val))
        text = "\n".join(chunks)
        score, matched = score_text(text, path.name)

        if score >= min_score:
            candidate = safe_copy(path, out_candidates, "xlsx__")
            return [
                Finding(
                    source_file=str(path),
                    file_type="xlsx",
                    page_or_item="workbook_text",
                    score=score,
                    matched_terms=", ".join(matched),
                    candidate_artifact=str(candidate),
                    notes="XLSX text matched",
                )
            ]
    except Exception as e:
        return [
            Finding(
                source_file=str(path),
                file_type="xlsx",
                page_or_item="error",
                score=0,
                matched_terms="",
                candidate_artifact="",
                notes=f"Cannot read XLSX: {e}",
            )
        ]

    return []


def process_zip(
    path: Path,
    out_candidates: Path,
    min_score: int,
    ocr_lang: str,
    tesseract_ok: bool,
    pdf_max_pages: Optional[int],
    max_zip_members: int,
) -> list[Finding]:
    findings: list[Finding] = []

    try:
        with zipfile.ZipFile(path, "r") as z:
            members = [
                m for m in z.namelist()
                if not m.endswith("/")
                and Path(m).suffix.lower() in (IMAGE_EXTS | PDF_EXTS | DOCX_EXTS | XLSX_EXTS)
            ][:max_zip_members]

            with tempfile.TemporaryDirectory() as td:
                tmp_dir = Path(td)
                for member in members:
                    try:
                        extracted = z.extract(member, tmp_dir)
                        extracted_path = Path(extracted)
                        inner_findings = process_file(
                            extracted_path,
                            out_candidates,
                            min_score,
                            ocr_lang,
                            tesseract_ok,
                            pdf_max_pages,
                            include_archives=False,
                            max_zip_members=max_zip_members,
                        )
                        for f in inner_findings:
                            f.source_file = f"{path}::{member}"
                            f.notes = "inside ZIP; " + f.notes
                            findings.append(f)
                    except Exception as e:
                        findings.append(
                            Finding(
                                source_file=f"{path}::{member}",
                                file_type="zip_member",
                                page_or_item="error",
                                score=0,
                                matched_terms="",
                                candidate_artifact="",
                                notes=f"Cannot extract/process ZIP member: {e}",
                            )
                        )
    except Exception as e:
        findings.append(
            Finding(
                source_file=str(path),
                file_type="zip",
                page_or_item="error",
                score=0,
                matched_terms="",
                candidate_artifact="",
                notes=f"Cannot open ZIP: {e}",
            )
        )

    return findings


def process_file(
    path: Path,
    out_candidates: Path,
    min_score: int,
    ocr_lang: str,
    tesseract_ok: bool,
    pdf_max_pages: Optional[int],
    include_archives: bool,
    max_zip_members: int,
) -> list[Finding]:
    suffix = path.suffix.lower()

    if suffix in IMAGE_EXTS:
        return process_image(path, out_candidates, min_score, ocr_lang, tesseract_ok)

    if suffix in PDF_EXTS:
        return process_pdf(path, out_candidates, min_score, ocr_lang, tesseract_ok, pdf_max_pages)

    if suffix in DOCX_EXTS:
        return process_docx(path, out_candidates, min_score)

    if suffix in XLSX_EXTS:
        return process_xlsx(path, out_candidates, min_score)

    if include_archives and suffix in ZIP_EXTS:
        return process_zip(
            path,
            out_candidates,
            min_score,
            ocr_lang,
            tesseract_ok,
            pdf_max_pages,
            max_zip_members,
        )

    return []


def write_report(report_path: Path, findings: list[Finding]) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = list(asdict(Finding("", "", "", 0, "", "", "")).keys())

    with report_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()

        for finding in sorted(findings, key=lambda x: x.score, reverse=True):
            writer.writerow(asdict(finding))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Find files/pages that likely contain a land plot / site plan."
    )
    parser.add_argument("root", help="Root folder to scan")
    parser.add_argument("--out", default="site_plan_search_output", help="Output folder")
    parser.add_argument("--min-score", type=int, default=5, help="Minimum score to mark as candidate")
    parser.add_argument(
        "--ocr-lang",
        default="rus+eng",
        help="Tesseract languages, e.g. rus+eng or eng",
    )
    parser.add_argument(
        "--pdf-max-pages",
        type=int,
        default=None,
        help="Process only first N pages of each PDF. Useful for testing.",
    )
    parser.add_argument(
        "--no-archives",
        action="store_true",
        help="Do not scan ZIP archives.",
    )
    parser.add_argument(
        "--max-zip-members",
        type=int,
        default=200,
        help="Max files to inspect inside one ZIP.",
    )

    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    out = Path(args.out).expanduser().resolve()
    out_candidates = out / "candidates"
    report_path = out / "report.csv"

    if not root.exists() or not root.is_dir():
        print(f"ERROR: root folder does not exist or is not a directory: {root}", file=sys.stderr)
        return 2

    out.mkdir(parents=True, exist_ok=True)
    out_candidates.mkdir(parents=True, exist_ok=True)

    tesseract_ok = ensure_tesseract_configured()

    print(f"Root: {root}")
    print(f"Output: {out}")
    print(f"Minimum score: {args.min_score}")
    print(f"Tesseract OCR: {'available' if tesseract_ok else 'NOT available'}")
    if not tesseract_ok:
        print(
            "WARNING: OCR is unavailable. The script will still search filenames and text layers, "
            "but scanned PDFs/images will be much harder to find.",
            file=sys.stderr,
        )

    files = list(iter_files(root, include_archives=not args.no_archives))
    print(f"Files to inspect: {len(files)}")

    findings: list[Finding] = []

    iterator = tqdm(files, desc="Scanning") if tqdm is not None else files

    for path in iterator:
        try:
            findings.extend(
                process_file(
                    path=path,
                    out_candidates=out_candidates,
                    min_score=args.min_score,
                    ocr_lang=args.ocr_lang,
                    tesseract_ok=tesseract_ok,
                    pdf_max_pages=args.pdf_max_pages,
                    include_archives=not args.no_archives,
                    max_zip_members=args.max_zip_members,
                )
            )
        except KeyboardInterrupt:
            print("\nInterrupted by user.", file=sys.stderr)
            break
        except Exception as e:
            findings.append(
                Finding(
                    source_file=str(path),
                    file_type=path.suffix.lower().lstrip("."),
                    page_or_item="error",
                    score=0,
                    matched_terms="",
                    candidate_artifact="",
                    notes=f"Unhandled error: {e}",
                )
            )

    write_report(report_path, findings)

    print()
    print(f"Done.")
    print(f"Findings: {len([f for f in findings if f.score >= args.min_score])}")
    print(f"Report: {report_path}")
    print(f"Candidate previews/files: {out_candidates}")

    print()
    print("Next steps:")
    print("1. Open report.csv in Excel.")
    print("2. Sort by score descending.")
    print("3. Open files from candidate_artifact.")
    print("4. If too many false positives: rerun with --min-score 8 or 10.")
    print("5. If too few results: rerun with --min-score 2 or 3.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
